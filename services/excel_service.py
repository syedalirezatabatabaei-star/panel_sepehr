from pathlib import Path
from openpyxl import load_workbook, Workbook


class ExcelService:

    def __init__(self, input_file: Path):
        self.input_file = Path(input_file)

    def read_contacts(self):

        if not self.input_file.exists():
            raise FileNotFoundError(
                f"Excel file not found: {self.input_file}"
            )

        workbook = load_workbook(
            self.input_file,
            read_only=True,
            data_only=True
        )

        sheet = workbook.active

        rows = list(sheet.iter_rows(values_only=True))

        workbook.close()

        if not rows:
            return []

        headers = [
            str(value).strip().lower()
            if value is not None
            else ""
            for value in rows[0]
        ]

        if "phone" not in headers:
            raise ValueError(
                "Excel must contain a 'phone' column."
            )

        if "name" not in headers:
            raise ValueError(
                "Excel must contain a 'name' column."
            )

        phone_index = headers.index("phone")
        name_index = headers.index("name")

        contacts = []

        for row_number, row in enumerate(rows[1:], start=2):

            if not row:
                continue

            phone = (
                row[phone_index]
                if phone_index < len(row)
                else None
            )

            name = (
                row[name_index]
                if name_index < len(row)
                else None
            )

            if phone is None:
                continue

            phone = str(phone).strip()

            if not phone:
                continue

            if name is None:
                name = ""

            name = str(name).strip()

            contacts.append({
                "row": row_number,
                "phone": phone,
                "name": name
            })

        return contacts

    def create_report(self, report_file: Path, results):

        report_file = Path(report_file)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Report"

        headers = [
            "phone",
            "name",
            "status",
            "message"
        ]

        sheet.append(headers)

        for result in results:

            sheet.append([
                result.get("phone", ""),
                result.get("name", ""),
                result.get("status", ""),
                result.get("message", "")
            ])

        # تنظیم عرض ستون‌ها
        sheet.column_dimensions["A"].width = 20
        sheet.column_dimensions["B"].width = 25
        sheet.column_dimensions["C"].width = 20
        sheet.column_dimensions["D"].width = 50

        workbook.save(report_file)

        return report_file

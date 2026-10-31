"""Authenticated Span API adapter for three isolated GoOmni provider devices."""
import base64
import json
import logging
import mimetypes
import os
import re
import threading
import time
import uuid
from datetime import datetime
from pathlib import Path

import requests
from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

router = APIRouter(prefix="/api/span", tags=["Span API"])
GOOMNI_URL = os.getenv("GOOMNI_URL", "http://eita-rubika-bale:3000").rstrip("/")
GOOMNI_AUTH = os.getenv("GOOMNI_BASIC_AUTH", "")
STATE = Path(os.getenv("SPAN_STATE_DIR", "/app/data/span-api"))
MEDIA, REPORTS, JOBS_DIR = STATE / "media", STATE / "reports", STATE / "jobs"
for folder in (STATE, MEDIA, REPORTS, JOBS_DIR):
    folder.mkdir(parents=True, exist_ok=True)

DEVICES = {
    "bale": os.getenv("SPAN_BALE_DEVICE_ID", "sepehr-span-bale"),
    "eitaa": os.getenv("SPAN_EITAA_DEVICE_ID", "sepehr-span-eitaa"),
    "rubika": os.getenv("SPAN_RUBIKA_DEVICE_ID", "sepehr-span-rubika"),
}
MAX_EXCEL = 10 * 1024 * 1024
MAX_IMAGE = 20 * 1024 * 1024
MAX_VIDEO = 20 * 1024 * 1024
MAX_FILE = 20 * 1024 * 1024
MAX_RECIPIENTS = 5000
JOBS: dict = {}
LOCK = threading.RLock()
LOG = logging.getLogger(__name__)


class GatewayOperationError(Exception):
    def __init__(self, message, result):
        super().__init__(message)
        self.result = result


def provider(value):
    key = str(value or "").lower()
    if key not in DEVICES:
        raise HTTPException(404, "پیام‌رسان Span API معتبر نیست")
    return key


def _device(p):
    return DEVICES[p]


def _account_path(p):
    return STATE / (p + "-account.json")


def _read_account(p):
    try:
        return json.loads(_account_path(p).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def _write_json(path, data):
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    try:
        tmp.chmod(0o600)
    except OSError:
        pass
    tmp.replace(path)


def _write_account(p, data):
    _write_json(_account_path(p), data)


def _headers(p, instance_id=None, idempotency_key=None):
    if not GOOMNI_AUTH or ":" not in GOOMNI_AUTH:
        raise HTTPException(503, "GOOMNI_BASIC_AUTH روی سرور تنظیم نشده است")
    headers = {
        "Authorization": "Basic " + base64.b64encode(GOOMNI_AUTH.encode()).decode(),
        "Accept": "application/json",
    }
    if instance_id:
        headers.update({"X-Device-Id": _device(p), "X-Device-Instance": str(instance_id)})
    if idempotency_key:
        headers["Idempotency-Key"] = str(idempotency_key)
    return headers


def _request(p, method, path, *, instance_id=None, idempotency_key=None, expected=(200,), **kwargs):
    try:
        response = requests.request(
            method, GOOMNI_URL + path,
            headers={**_headers(p, instance_id, idempotency_key), **kwargs.pop("headers", {})},
            timeout=(10, 90), **kwargs
        )
        try:
            payload = response.json()
        except ValueError:
            payload = {"message": response.text[:500]}
    except requests.RequestException as exc:
        LOG.warning("GoOmni transport failure (%s)", type(exc).__name__)
        raise HTTPException(503, "ارتباط با GoOmni برقرار نشد؛ نتیجه عملیات احتمالا نیازمند بررسی است")
    result = payload.get("results", payload.get("Results", payload)) if isinstance(payload, dict) else payload
    if response.status_code not in expected:
        message = payload.get("message") or payload.get("detail") or payload.get("code") or "خطای GoOmni"
        if isinstance(result, dict) and result.get("send_id"):
            raise GatewayOperationError(str(message), result)
        raise HTTPException(response.status_code, str(message)[:500])
    return result


def _ensure_device(p):
    alias = _device(p)
    devices = _request(p, "GET", "/devices")
    items = devices if isinstance(devices, list) else devices.get("devices", []) if isinstance(devices, dict) else []
    found = next((x for x in items if str(x.get("id") or x.get("device_id")) == alias), None)
    if found:
        actual = str(found.get("provider") or "").lower()
        if actual and actual != p:
            raise HTTPException(409, "این شناسه دستگاه در GoOmni متعلق به پیام‌رسان دیگری است")
        return found
    return _request(
        p, "POST", "/devices",
        json={"device_id": alias, "provider": p},
        idempotency_key="sepehr-span-provision-" + p + "-v1",
        expected=(200, 201),
    )


def _instance(p):
    saved = _read_account(p)
    if saved.get("instance_id") and saved.get("device_id") == _device(p):
        return str(saved["instance_id"])
    device_data = _ensure_device(p)
    value = device_data.get("instance_id") or device_data.get("instanceId")
    if not value:
        raise HTTPException(503, "GoOmni شناسه instance دستگاه را برنگرداند")
    saved.update({"provider": p, "device_id": _device(p), "instance_id": str(value)})
    _write_account(p, saved)
    return str(value)


def _status(p):
    instance_id = _instance(p)
    result = _request(p, "GET", f"/devices/{_device(p)}/status", instance_id=instance_id)
    state = str(result.get("auth") or "").lower()
    saved = _read_account(p)
    return {
        "provider": p,
        "device_id": _device(p),
        "exists": state in {"authenticated", "authorized", "ready"},
        "state": state or "unknown",
        "phone": saved.get("phone", ""),
        "challenge": result.get("challenge"),
    }


@router.get("/{provider_name}/account")
def account_status(provider_name: str):
    return _status(provider(provider_name))


@router.post("/{provider_name}/auth/send-code")
def send_code(provider_name: str, payload: dict):
    p = provider(provider_name)
    phone = str(payload.get("phone") or "").strip()
    if not re.fullmatch(r"\+?[0-9۰-۹٠-٩]{8,16}", phone):
        raise HTTPException(422, "شماره تلفن معتبر نیست")
    instance_id = _instance(p)
    result = _request(
        p, "POST", f"/devices/{_device(p)}/login",
        instance_id=instance_id, json={"phone": phone}
    )
    saved = _read_account(p)
    saved.update({
        "provider": p, "device_id": _device(p), "instance_id": instance_id,
        "phone": phone, "challenge_id": result.get("challenge_id") or result.get("id"),
    })
    _write_account(p, saved)
    return {
        "success": True, "phone": phone, "challenge_id": saved.get("challenge_id"),
        "delivery": result.get("delivery"), "next_delivery": result.get("next_delivery"),
        "message": result.get("message") or "درخواست کد ورود ثبت شد",
    }


@router.post("/{provider_name}/auth/login")
def login(provider_name: str, payload: dict):
    p = provider(provider_name)
    saved = _read_account(p)
    challenge_id = str(payload.get("challenge_id") or saved.get("challenge_id") or "")
    code = str(payload.get("code") or "").strip()
    if not challenge_id or not code:
        raise HTTPException(422, "کد ورود یا شناسه چالش معتبر نیست")
    result = _request(
        p, "POST", f"/devices/{_device(p)}/login/code",
        instance_id=_instance(p), json={"challenge_id": challenge_id, "code": code}
    )
    state = str(result.get("state") or "").lower()
    if result.get("requires_password") or state == "awaiting_password":
        return {"success": False, "requires_password": True, "challenge_id": challenge_id}
    status = _request(p, "GET", f"/devices/{_device(p)}/status", instance_id=_instance(p))
    saved.pop("challenge_id", None)
    saved["account_id"] = status.get("account_id")
    _write_account(p, saved)
    return {"success": True, "phone": saved.get("phone", ""), "state": status.get("auth") or state or "authenticated"}


@router.post("/{provider_name}/auth/login-password")
def login_password(provider_name: str, payload: dict):
    p = provider(provider_name)
    saved = _read_account(p)
    challenge_id = str(payload.get("challenge_id") or saved.get("challenge_id") or "")
    password = str(payload.get("password") or "")
    if not challenge_id or not password:
        raise HTTPException(422, "شناسه چالش و رمز دومرحله‌ای الزامی است")
    result = _request(
        p, "POST", f"/devices/{_device(p)}/login/password",
        instance_id=_instance(p), json={"challenge_id": challenge_id, "password": password}
    )
    status = _request(p, "GET", f"/devices/{_device(p)}/status", instance_id=_instance(p))
    saved.pop("challenge_id", None)
    saved["account_id"] = status.get("account_id")
    _write_account(p, saved)
    return {"success": True, "phone": saved.get("phone", ""), "state": status.get("auth") or result.get("state") or "authenticated"}


async def _save_upload(upload, path, limit, label):
    size = 0
    with path.open("wb") as output:
        while True:
            chunk = await upload.read(1024 * 1024)
            if not chunk:
                break
            size += len(chunk)
            if size > limit:
                output.close()
                path.unlink(missing_ok=True)
                raise HTTPException(413, f"حجم {label} بیش از حد مجاز است")
            output.write(chunk)
    if not size:
        path.unlink(missing_ok=True)
        raise HTTPException(400, f"فایل {label} خالی است")
    return path


def _recipients(path, p):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        headers = [str(c.value or "").strip().lower() for c in sheet[1]]
        peer_headers = {"peer_id", "user_id", "guid", "شناسه", "شناسه کاربری", "شناسه مخاطب"}
        name_headers = {"name", "نام", "نام مخاطب"}
        type_headers = {"peer_type", "type", "نوع", "نوع مخاطب"}
        peer_col = next((i + 1 for i, x in enumerate(headers) if x in peer_headers), 1)
        name_col = next((i + 1 for i, x in enumerate(headers) if x in name_headers), None)
        type_col = next((i + 1 for i, x in enumerate(headers) if x in type_headers), None)
        has_header = any(x in peer_headers for x in headers)
        targets, seen = [], set()
        for row in sheet.iter_rows(min_row=2 if has_header else 1, values_only=True):
            raw = row[peer_col - 1] if len(row) >= peer_col else None
            peer_id = str(raw or "").strip()
            if not peer_id or len(peer_id) > 256:
                continue
            peer_type = str(row[type_col - 1] if type_col and len(row) >= type_col else "user").strip().lower()
            if peer_type not in {"user", "group", "channel"}:
                continue
            if p == "rubika" and peer_type != "user":
                raise ValueError("در نسخه فعلی Span API، ارسال روبیکا فقط به مخاطب user پشتیبانی می‌شود")
            if peer_type == "user" and peer_id.startswith("channel_"):
                raise ValueError("شناسه channel_ نباید با نوع user ارسال شود")
            name = str(row[name_col - 1] or "").strip() if name_col and len(row) >= name_col else ""
            key = (peer_type, peer_id)
            if key not in seen:
                seen.add(key)
                targets.append({"peer": {"type": peer_type, "id": peer_id}, "name": name or peer_id})
            if len(targets) > MAX_RECIPIENTS:
                raise ValueError(f"حداکثر {MAX_RECIPIENTS} مخاطب در هر ارسال مجاز است")
        return targets
    finally:
        workbook.close()


def _job_path(job_id):
    return JOBS_DIR / (job_id + ".json")


def _save_job(job):
    _write_json(_job_path(job["job_id"]), job)


def _get_job(job_id):
    with LOCK:
        if job_id in JOBS:
            return JOBS[job_id]
        try:
            job = json.loads(_job_path(job_id).read_text(encoding="utf-8"))
            JOBS[job_id] = job
            return job
        except (OSError, ValueError):
            return None


def _send_one(p, instance_id, target, message, media_path, media_type, key):
    peer = target["peer"]
    if media_path:
        mime = mimetypes.guess_type(media_path.name)[0] or "application/octet-stream"
        with media_path.open("rb") as stream:
            result = _request(
                p, "POST", "/send/" + media_type, instance_id=instance_id,
                idempotency_key=key, expected=(200, 202),
                data={"peer": json.dumps(peer, ensure_ascii=False), "caption": message},
                files={media_type: (media_path.name, stream, mime)},
            )
    else:
        result = _request(
            p, "POST", "/send/message", instance_id=instance_id,
            idempotency_key=key, expected=(200, 202),
            json={"peer": peer, "message": message},
        )
    send_id = result.get("send_id")
    state = str(result.get("state") or "").lower()
    if state in {"queued", "sending"} and send_id:
        deadline = time.monotonic() + 90
        while time.monotonic() < deadline:
            time.sleep(2)
            operation = _request(p, "GET", "/send/operations/" + str(send_id), instance_id=instance_id)
            state = str(operation.get("state") or "").lower()
            if state not in {"queued", "sending"}:
                result = operation
                break
    if state == "succeeded" or result.get("message_id"):
        return {"status": "success", "send_id": send_id, "message_id": result.get("message_id"), "detail": "Gateway عملیات را موفق اعلام کرد"}
    if state == "unknown":
        return {"status": "unknown", "send_id": send_id, "detail": "نتیجه ارسال نامشخص است؛ ارسال خودکار تکرار نشد"}
    return {"status": "error", "send_id": send_id, "detail": str(result.get("error_message") or result.get("message") or state or "نتیجه ارسال تأیید نشد")[:500]}


def _report(job):
    path = REPORTS / ("span_" + job["job_id"] + ".xlsx")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "نتایج"
    sheet.append(["ردیف", "پیام‌رسان", "شناسه مخاطب", "نوع مخاطب", "نام", "وضعیت", "شناسه عملیات", "توضیحات"])
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for index, row in enumerate(job["results"], 1):
        sheet.append([index, job["provider"], row["peer_id"], row["peer_type"], row["name"], row["status"], row.get("send_id") or "", row.get("detail") or ""])
    for col, width in {"A": 8, "B": 14, "C": 34, "D": 14, "E": 24, "F": 18, "G": 38, "H": 60}.items():
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    summary = workbook.create_sheet("خلاصه")
    summary.append(["شاخص", "مقدار"])
    for label, key in [("کل", "total"), ("موفق", "success"), ("ناموفق", "failed"), ("نامشخص", "unknown")]:
        summary.append([label, job[key]])
    workbook.save(path)
    return str(path)


def _run_job(job_id, targets, message, media_path, media_type, delay):
    job = _get_job(job_id)
    job["status"], job["phase"] = "running", "در حال ارسال"
    _save_job(job)
    results = []
    for index, target in enumerate(targets, 1):
        peer = target["peer"]
        row = {"peer_id": peer["id"], "peer_type": peer["type"], "name": target["name"], "status": "error"}
        try:
            row.update(_send_one(job["provider"], job["instance_id"], target, message, media_path, media_type,
                                 f"sepehr-span-{job_id}-{index}"))
        except GatewayOperationError as exc:
            row.update({"status": "unknown" if str(exc.result.get("state")) == "unknown" else "error",
                        "send_id": exc.result.get("send_id"), "detail": str(exc)[:500]})
        except HTTPException as exc:
            row.update({"status": "unknown" if exc.status_code == 503 else "error", "detail": str(exc.detail)[:500]})
        except Exception as exc:
            LOG.exception("Span send failed")
            row.update({"status": "unknown", "detail": type(exc).__name__})
        results.append(row)
        job["results"] = results
        job["processed"] = index
        job["success"] = sum(x["status"] == "success" for x in results)
        job["failed"] = sum(x["status"] == "error" for x in results)
        job["unknown"] = sum(x["status"] == "unknown" for x in results)
        job["progress"] = int(index * 100 / max(1, len(targets)))
        job["phase"] = "در حال پردازش مخاطبان" if index < len(targets) else "در حال ساخت گزارش"
        _save_job(job)
        if index < len(targets) and delay:
            time.sleep(delay)
    try:
        job["report_path"] = _report(job)
    except Exception:
        LOG.exception("Span report creation failed")
        job["report_path"] = None
    job.update({"status": "completed", "progress": 100, "phase": "پایان یافت",
                "finished_at": datetime.now().isoformat(timespec="seconds")})
    _save_job(job)
    if media_path:
        media_path.unlink(missing_ok=True)


@router.post("/{provider_name}/send")
async def start_send(provider_name: str, excel: UploadFile = File(...), message: str = Form(""),
                     media: UploadFile | None = File(None), media_type: str = Form("text"),
                     delay_seconds: int = Form(0, ge=0, le=30)):
    p = provider(provider_name)
    instance_id = _instance(p)
    if not _status(p)["exists"]:
        raise HTTPException(401, "ابتدا حساب این پیام‌رسان را در Span API متصل کنید")
    if not excel.filename or Path(excel.filename).suffix.lower() not in {".xlsx", ".xlsm"}:
        raise HTTPException(400, "فایل مخاطبان باید xlsx یا xlsm باشد")
    if media_type not in {"text", "image", "video", "file"}:
        raise HTTPException(422, "نوع رسانه نامعتبر است")
    if media_type == "text" and not message.strip():
        raise HTTPException(400, "متن پیام الزامی است")
    if media_type != "text" and (not media or not media.filename):
        raise HTTPException(400, "برای نوع پیام انتخاب‌شده فایل لازم است")
    job_id = str(uuid.uuid4())
    excel_path = MEDIA / (job_id + "_contacts.xlsx")
    media_path = None
    try:
        await _save_upload(excel, excel_path, MAX_EXCEL, "Excel")
        try:
            targets = _recipients(excel_path, p)
        except ValueError as exc:
            raise HTTPException(400, str(exc))
        if not targets:
            raise HTTPException(400, "هیچ مخاطب معتبری با ستون peer_id در فایل پیدا نشد")
        if media_type != "text":
            suffix = Path(media.filename).suffix.lower()
            allowed = {
                "image": {".jpg", ".jpeg", ".png", ".webp"},
                "video": {".mp4", ".mov", ".m4v"},
                "file": {".pdf", ".doc", ".docx", ".zip"},
            }[media_type]
            if suffix not in allowed:
                raise HTTPException(400, "پسوند فایل با نوع رسانه مطابقت ندارد")
            limit = MAX_IMAGE if media_type == "image" else MAX_VIDEO if media_type == "video" else MAX_FILE
            media_path = await _save_upload(media, MEDIA / (job_id + "_media" + suffix), limit, "رسانه")
        job = {
            "job_id": job_id, "provider": p, "instance_id": instance_id,
            "status": "queued", "phase": "در صف انتظار", "total": len(targets),
            "processed": 0, "success": 0, "failed": 0, "unknown": 0,
            "progress": 0, "results": [], "created_at": datetime.now().isoformat(timespec="seconds"),
        }
        JOBS[job_id] = job
        _save_job(job)
        threading.Thread(target=_run_job, args=(job_id, targets, message, media_path, media_type, delay_seconds), daemon=True).start()
        return {"success": True, "job_id": job_id, "total": len(targets)}
    except Exception:
        excel_path.unlink(missing_ok=True)
        if media_path:
            media_path.unlink(missing_ok=True)
        raise
    finally:
        excel_path.unlink(missing_ok=True)


@router.get("/{provider_name}/send/status/{job_id}")
def send_status(provider_name: str, job_id: str):
    p = provider(provider_name)
    job = _get_job(job_id)
    if not job or job.get("provider") != p:
        raise HTTPException(404, "شناسه ارسال پیدا نشد")
    return {key: value for key, value in job.items() if key not in {"report_path", "instance_id"}}


@router.get("/{provider_name}/send/export/{job_id}")
def export_report(provider_name: str, job_id: str):
    p = provider(provider_name)
    job = _get_job(job_id)
    if not job or job.get("provider") != p or not job.get("report_path"):
        raise HTTPException(404, "گزارش هنوز آماده نیست")
    path = Path(job["report_path"])
    if not path.is_file() or path.parent != REPORTS:
        raise HTTPException(404, "گزارش پیدا نشد")
    return FileResponse(path, filename=path.name,
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

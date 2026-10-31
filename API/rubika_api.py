import json
import logging
import mimetypes
import os
import threading
import time
import uuid
import unicodedata
from datetime import datetime
from pathlib import Path

import requests
from fastapi import APIRouter, BackgroundTasks, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

router = APIRouter(prefix="/api/rubika", tags=["Rubika personal messaging"])
GOOMNI_URL = os.getenv("GOOMNI_URL", "http://goomni:3000").rstrip("/")
GOOMNI_AUTH = os.getenv("GOOMNI_BASIC_AUTH", "")
DEVICE_ID = os.getenv("RUBIKA_DEVICE_ID", "sepehr-rubika")
STATE_DIR = Path(os.getenv("RUBIKA_STATE_DIR", "/app/data/rubika-state"))
MEDIA_DIR = STATE_DIR / "media"
REPORT_DIR = STATE_DIR / "reports"
ACCOUNT_FILE = STATE_DIR / "account.json"
for directory in (STATE_DIR, MEDIA_DIR, REPORT_DIR):
    directory.mkdir(parents=True, exist_ok=True)

JOBS = {}
JOBS_LOCK = threading.Lock()
HTTP_TIMEOUT = (10, 90)
logger = logging.getLogger(__name__)


def _headers(instance_id=None, idempotency_key=None):
    if not GOOMNI_AUTH or ":" not in GOOMNI_AUTH:
        raise HTTPException(503, "تنظیم GOOMNI_BASIC_AUTH روی سرور انجام نشده است")
    headers = {"Authorization": "Basic " + __import__("base64").b64encode(GOOMNI_AUTH.encode()).decode(),
               "Accept": "application/json"}
    if instance_id:
        headers.update({"X-Device-Id": DEVICE_ID, "X-Device-Instance": str(instance_id)})
    if idempotency_key:
        headers["Idempotency-Key"] = idempotency_key
    return headers


def _request(method, path, *, instance_id=None, idempotency_key=None, **kwargs):
    try:
        response = requests.request(method, GOOMNI_URL + path, headers={**_headers(instance_id, idempotency_key), **kwargs.pop("headers", {})},
                                    timeout=HTTP_TIMEOUT, **kwargs)
        try:
            payload = response.json()
        except ValueError:
            payload = {"message": response.text[:500]}
        if response.status_code >= 400:
            message = payload.get("message") or payload.get("detail") or payload.get("Message") or payload.get("code") or "خطا در Gateway روبیکا"
            raise HTTPException(response.status_code, str(message)[:500])
        return payload.get("results", payload.get("Results", payload)) if isinstance(payload, dict) else payload
    except requests.RequestException as exc:
        logger.warning("GoOmni unavailable: %s", type(exc).__name__)
        raise HTTPException(503, "ارتباط با سرویس GoOmni برقرار نشد")


def _account():
    try:
        return json.loads(ACCOUNT_FILE.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def _save_account(data):
    ACCOUNT_FILE.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    try:
        ACCOUNT_FILE.chmod(0o600)
    except OSError:
        pass


def _devices():
    result = _request("GET", "/devices")
    return result if isinstance(result, list) else result.get("devices", []) if isinstance(result, dict) else []


def _ensure_device():
    devices = _devices()
    device = next((item for item in devices if str(item.get("id") or item.get("device_id")) == DEVICE_ID), None)
    if device:
        return device
    return _request("POST", "/devices", json={"provider": "rubika", "device_id": DEVICE_ID},
                    idempotency_key="sepehr-rubika-device-v1")


def _instance_id(device=None):
    device = device or next((x for x in _devices() if str(x.get("id") or x.get("device_id")) == DEVICE_ID), None)
    if not device:
        raise HTTPException(503, "دستگاه روبیکا در Gateway ساخته نشده است")
    value = device.get("instance_id") or device.get("instanceId")
    if not value:
        raise HTTPException(503, "شناسه نشست Gateway دریافت نشد؛ فهرست دستگاه‌ها را بررسی کنید")
    return str(value)


def _phone(value):
    digits = []
    for char in str(value or ""):
        try:
            digits.append(str(unicodedata.digit(char)))
        except (TypeError, ValueError):
            if char.isascii() and char.isdigit():
                digits.append(char)
    value = "".join(digits)
    if value.startswith("00"):
        value = value[2:]
    if value.startswith("0") and len(value) == 11:
        value = "+98" + value[1:]
    elif value.startswith("9") and len(value) == 10:
        value = "+98" + value
    elif value.startswith("98") and len(value) == 12:
        value = "+" + value
    return value if value.startswith("+") and 10 <= len(value) <= 16 else ""


@router.get("/health")
def health():
    try:
        result = _request("GET", "/ready")
        return {"success": True, "service": "GoOmni", "status": "online", "ready": result}
    except HTTPException as exc:
        raise HTTPException(503, exc.detail)


@router.get("/account")
def account_status():
    saved = _account()
    try:
        device = next((x for x in _devices() if str(x.get("id") or x.get("device_id")) == DEVICE_ID), None)
        if not device:
            return {"exists": False, "configured": True}
        instance = _instance_id(device)
        status = _request("GET", f"/devices/{DEVICE_ID}/status", instance_id=instance)
        login_state = _request("GET", f"/devices/{DEVICE_ID}/login", instance_id=instance)
        auth_state = str((status or {}).get("auth") or (login_state or {}).get("state") or "").lower()
        authenticated = auth_state in {"authenticated", "authorized", "ready", "connected"}
        phone = (saved or {}).get("phone") or device.get("phone") or ""
        if authenticated:
            _save_account({"phone": phone, "device_id": DEVICE_ID, "instance_id": instance})
        return {"exists": authenticated, "phone": phone, "device_id": DEVICE_ID,
                "state": auth_state or "unknown", "status": status, "login": login_state}
    except HTTPException:
        raise


@router.post("/auth/send-code")
def send_code(payload: dict):
    phone = _phone(payload.get("phone"))
    if not phone:
        raise HTTPException(422, "شماره تلفن معتبر نیست")
    device = _ensure_device()
    instance = _instance_id(device)
    result = _request("POST", f"/devices/{DEVICE_ID}/login", instance_id=instance, json={"phone": phone})
    challenge_id = result.get("challenge_id") or result.get("id")
    state = result.get("state")
    if challenge_id:
        _save_account({"phone": phone, "device_id": DEVICE_ID, "instance_id": instance, "challenge_id": challenge_id})
    return {"success": True, "phone": phone, "challenge_id": challenge_id, "state": state,
            "message": result.get("message") or "درخواست ورود به Gateway ارسال شد"}


@router.post("/auth/login")
def login(payload: dict):
    saved = _account()
    code = str(payload.get("code") or "").strip()
    challenge_id = str(payload.get("challenge_id") or saved.get("challenge_id") or "")
    if not code or not challenge_id:
        raise HTTPException(422, "کد ورود یا شناسه چالش معتبر نیست")
    instance = _instance_id()
    result = _request("POST", f"/devices/{DEVICE_ID}/login/code", instance_id=instance,
                      json={"challenge_id": challenge_id, "code": code})
    state = str(result.get("state") or "").lower()
    if result.get("requires_password") or state in {"awaiting_password", "password_required"}:
        return {"success": False, "requires_password": True, "phone": saved.get("phone"), "challenge_id": challenge_id}
    _save_account({"phone": saved.get("phone", ""), "device_id": DEVICE_ID, "instance_id": instance})
    return {"success": True, "phone": saved.get("phone", ""), "state": state or "authenticated"}


@router.post("/auth/login-password")
def login_password(payload: dict):
    saved = _account()
    challenge_id = str(payload.get("challenge_id") or saved.get("challenge_id") or "")
    password = str(payload.get("password") or "")
    if not challenge_id or not password:
        raise HTTPException(422, "شناسه چالش و رمز دومرحله‌ای الزامی است")
    instance = _instance_id()
    result = _request("POST", f"/devices/{DEVICE_ID}/login/password", instance_id=instance,
                      json={"challenge_id": challenge_id, "password": password})
    _save_account({"phone": saved.get("phone", ""), "device_id": DEVICE_ID, "instance_id": instance})
    return {"success": True, "phone": saved.get("phone", ""), "state": result.get("state", "authenticated")}


@router.post("/account/reset")
def reset_account():
    saved = _account()
    if saved:
        try:
            _request("POST", f"/devices/{DEVICE_ID}/logout", instance_id=saved.get("instance_id") or _instance_id())
        except HTTPException:
            logger.info("Rubika logout was not confirmed; local account mapping will be cleared")
    ACCOUNT_FILE.unlink(missing_ok=True)
    return {"success": True}


def _load_recipients(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = [str(cell.value or "").strip().lower() for cell in ws[1]]
    phone_names = {"phone", "mobile", "number", "شماره", "شماره تلفن", "موبایل", "تلفن", "شماره همراه"}
    name_names = {"name", "firstname", "first name", "نام", "نام مخاطب"}
    phone_col = next((i + 1 for i, value in enumerate(headers) if value in phone_names), 1)
    name_col = next((i + 1 for i, value in enumerate(headers) if value in name_names), 2)
    start_row = 2 if any(value in phone_names for value in headers) else 1
    recipients, seen = [], set()
    for row in range(start_row, ws.max_row + 1):
        phone = _phone(ws.cell(row, phone_col).value)
        name = str(ws.cell(row, name_col).value or "").strip() if name_col <= ws.max_column else ""
        if phone and phone not in seen:
            seen.add(phone)
            recipients.append({"phone": phone, "name": name or phone})
    wb.close()
    return recipients


def _write_report(job_id, results):
    output = REPORT_DIR / ("rubika_report_" + job_id + ".xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "نتایج"
    headers = ["ردیف", "شماره", "نام", "وضعیت پیام", "توضیحات"]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for index, item in enumerate(results, 1):
        label = {"success": "موفق", "unknown": "نامشخص (ارسال را تکرار نکنید)", "error": "ناموفق"}.get(item.get("status"), "نامشخص")
        ws.append([index, item.get("phone", ""), item.get("name", ""), label, item.get("detail", "")])
    for col, width in {"A": 8, "B": 22, "C": 24, "D": 28, "E": 65}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    summary = wb.create_sheet("خلاصه")
    summary.append(["شاخص", "مقدار"])
    summary.append(["کل مخاطبین", len(results)])
    summary.append(["موفق", sum(1 for x in results if x.get("status") == "success")])
    summary.append(["ناموفق", sum(1 for x in results if x.get("status") == "error")])
    summary.append(["نامشخص", sum(1 for x in results if x.get("status") == "unknown")])
    wb.save(output)
    return str(output)


def _send_one(instance, target, message, media_path, media_kind):
    key = str(uuid.uuid4())
    if not media_path:
        result = _request("POST", "/send/message", instance_id=instance, idempotency_key=key,
                          json={"phone": target["phone"], "message": message})
    else:
        mime = mimetypes.guess_type(media_path.name)[0] or "application/octet-stream"
        field = media_kind
        with media_path.open("rb") as handle:
            result = _request("POST", "/send/" + media_kind, instance_id=instance, idempotency_key=key,
                             data={"phone": target["phone"], "caption": message},
                             files={field: (media_path.name, handle, mime)})
    send_id = result.get("send_id") or result.get("id")
    state = str(result.get("state") or result.get("status") or "").lower()
    if state in {"succeeded", "success", "message sent"}:
        return "success", "پیام توسط Gateway تأیید شد"
    if not send_id:
        return ("unknown", "پاسخ Gateway شناسه عملیات نداشت؛ برای جلوگیری از ارسال تکراری بررسی دستی لازم است")
    deadline = time.monotonic() + 45
    while time.monotonic() < deadline:
        operation = _request("GET", f"/send/operations/{send_id}", instance_id=instance)
        state = str(operation.get("state") or operation.get("status") or "").lower()
        if state == "succeeded" or state == "message sent":
            return "success", "پیام ارسال شد"
        if state in {"failed", "rejected"}:
            return "error", str(operation.get("error_message") or operation.get("error_code") or "Gateway ارسال را رد کرد")
        if state == "unknown":
            return "unknown", "نتیجه ارسال از سمت پیام‌رسان نامشخص است؛ ارسال خودکار تکرار نشد"
        time.sleep(1.5)
    return "unknown", "مهلت انتظار تمام شد؛ وضعیت عملیات را در Gateway بررسی کنید و ارسال را تکرار نکنید"


def _run_job(job_id, recipients, message, media_path, media_kind, delay, instance):
    results = []
    with JOBS_LOCK:
        JOBS[job_id]["status"] = "running"
        JOBS[job_id]["phase"] = "در حال ارسال پیام"
    for index, target in enumerate(recipients, 1):
        row = {"phone": target["phone"], "name": target["name"], "status": "error", "detail": ""}
        with JOBS_LOCK:
            JOBS[job_id]["current"] = target["phone"]
            JOBS[job_id]["processed"] = index - 1
            JOBS[job_id]["progress"] = int((index - 1) * 100 / max(1, len(recipients)))
            JOBS[job_id]["phase"] = "در حال ارسال به مخاطب " + str(index)
        try:
            row["status"], row["detail"] = _send_one(instance, target, message, media_path, media_kind)
        except HTTPException as exc:
            row["status"], row["detail"] = "error", str(exc.detail)[:500]
        except Exception as exc:
            row["status"], row["detail"] = "error", (type(exc).__name__ + ": " + str(exc))[:500]
        results.append(row)
        with JOBS_LOCK:
            JOBS[job_id]["processed"] = index
            JOBS[job_id]["progress"] = int(index * 100 / max(1, len(recipients)))
            JOBS[job_id]["success"] = sum(1 for x in results if x["status"] == "success")
            JOBS[job_id]["failed"] = sum(1 for x in results if x["status"] == "error")
            JOBS[job_id]["unknown"] = sum(1 for x in results if x["status"] == "unknown")
            JOBS[job_id]["results"] = list(results)
        if index < len(recipients) and delay:
            time.sleep(delay)
    try:
        report = _write_report(job_id, results)
    except Exception as exc:
        report = None
        logger.exception("Could not write Rubika report: %s", type(exc).__name__)
    with JOBS_LOCK:
        JOBS[job_id].update({"status": "completed", "phase": "پایان یافت", "progress": 100,
                             "finished_at": datetime.now().isoformat(timespec="seconds"), "report_path": report})
    if media_path:
        media_path.unlink(missing_ok=True)


@router.post("/send")
async def start_send(background_tasks: BackgroundTasks, excel: UploadFile = File(...), message: str = Form(""),
                     media: UploadFile | None = File(None), delay_seconds: int = Form(0, ge=0, le=30)):
    saved = _account()
    if not saved or not saved.get("instance_id"):
        raise HTTPException(401, "ابتدا حساب روبیکا را متصل کنید")
    if not message.strip() and not (media and media.filename):
        raise HTTPException(400, "متن پیام یا فایل رسانه‌ای الزامی است")
    if not excel.filename or Path(excel.filename).suffix.lower() not in {".xlsx", ".xlsm"}:
        raise HTTPException(400, "فایل مخاطبان باید Excel با پسوند xlsx یا xlsm باشد")
    job_id = str(uuid.uuid4())
    excel_path = MEDIA_DIR / (job_id + "_" + Path(excel.filename).name)
    media_path = None
    try:
        excel_path.write_bytes(await excel.read())
        recipients = _load_recipients(excel_path)
        if not recipients:
            raise HTTPException(400, "هیچ شماره معتبر و تکرارنشده‌ای در فایل پیدا نشد")
        if len(recipients) > 5000:
            raise HTTPException(400, "حداکثر ۵۰۰۰ مخاطب در هر نوبت مجاز است")
        media_kind = None
        if media and media.filename:
            suffix = Path(media.filename).suffix.lower()
            allowed = {".jpg": "image", ".jpeg": "image", ".png": "image", ".mp4": "video",
                       ".pdf": "file", ".doc": "file", ".docx": "file"}
            if suffix not in allowed:
                raise HTTPException(400, "فرمت رسانه پشتیبانی نمی‌شود")
            data = await media.read()
            if not data or len(data) > 64 * 1024 * 1024:
                raise HTTPException(400, "حجم رسانه خالی یا بیش از ۶۴ مگابایت است")
            if allowed[suffix] == "video" and len(data) > 20 * 1024 * 1024:
                raise HTTPException(400, "حجم ویدئو نباید بیشتر از ۲۰ مگابایت باشد")
            media_path = MEDIA_DIR / (job_id + "_media" + suffix)
            media_path.write_bytes(data)
            media_kind = allowed[suffix]
        with JOBS_LOCK:
            JOBS[job_id] = {"job_id": job_id, "status": "queued", "total": len(recipients),
                            "processed": 0, "success": 0, "failed": 0, "unknown": 0, "progress": 0,
                            "results": [], "created_at": datetime.now().isoformat(timespec="seconds")}
        thread = threading.Thread(target=_run_job, args=(job_id, recipients, message, media_path, media_kind,
                                                        int(delay_seconds), saved["instance_id"]), daemon=True)
        thread.start()
        return {"success": True, "job_id": job_id, "total": len(recipients)}
    except Exception:
        excel_path.unlink(missing_ok=True)
        if media_path:
            media_path.unlink(missing_ok=True)
        raise


@router.get("/send/status/{job_id}")
def send_status(job_id: str):
    with JOBS_LOCK:
        job = JOBS.get(job_id)
        if not job:
            raise HTTPException(404, "شناسه ارسال پیدا نشد")
        return {key: value for key, value in job.items() if key != "report_path"}


@router.get("/send/export/{job_id}")
def export_report(job_id: str):
    with JOBS_LOCK:
        job = JOBS.get(job_id)
        report = job.get("report_path") if job else None
    if not report or not Path(report).exists():
        raise HTTPException(404, "گزارش هنوز آماده نیست")
    return FileResponse(report, filename=Path(report).name,
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

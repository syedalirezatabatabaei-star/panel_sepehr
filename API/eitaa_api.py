import json
import os
import threading
import time
import uuid
from datetime import datetime
from pathlib import Path

import requests
from fastapi import APIRouter, BackgroundTasks, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

router = APIRouter(prefix="/api/eitaa", tags=["Eitaa broadcast"])
SERVICE_URL = os.getenv("EITAA_SERVICE_URL", "http://127.0.0.1:1234").rstrip("/")
STATE_DIR = Path(os.getenv("EITAA_STATE_DIR", "/app/data/eitaa-state"))
MEDIA_DIR = Path(os.getenv("EITAA_MEDIA_DIR", "/shared/eitaa-media"))
REPORT_DIR = STATE_DIR / "reports"
ACCOUNT_FILE = STATE_DIR / "account.json"
for directory in (STATE_DIR, MEDIA_DIR, REPORT_DIR):
    directory.mkdir(parents=True, exist_ok=True)
JOBS = {}
JOBS_LOCK = threading.Lock()
HTTP_TIMEOUT = (10, 120)


def _request(method, path, **kwargs):
    try:
        response = requests.request(method, SERVICE_URL + path, timeout=HTTP_TIMEOUT, **kwargs)
        try:
            payload = response.json()
        except ValueError:
            payload = {"detail": response.text[:1000]}
        if response.status_code >= 400:
            detail = payload.get("error") or payload.get("detail") or payload.get("message") or str(payload)
            raise HTTPException(response.status_code, str(detail))
        if isinstance(payload, dict) and payload.get("error"):
            raise HTTPException(400, str(payload["error"]))
        return payload
    except requests.RequestException as exc:
        raise HTTPException(503, "سرویس ایتا در دسترس نیست: " + str(exc))


def _account():
    if not ACCOUNT_FILE.exists():
        return None
    try:
        return json.loads(ACCOUNT_FILE.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None


def _save_account(phone):
    ACCOUNT_FILE.write_text(json.dumps({"phone": phone}, ensure_ascii=False), encoding="utf-8")


def _phone(value):
    digits = "".join(ch for ch in str(value or "") if ch.isdigit())
    if digits.startswith("00"):
        digits = digits[2:]
    if digits.startswith("0") and len(digits) == 11:
        digits = "98" + digits[1:]
    elif len(digits) == 10 and digits.startswith("9"):
        digits = "98" + digits
    return "+" + digits if 10 <= len(digits) <= 15 else ""


@router.get("/health")
def health():
    try:
        response = requests.get(SERVICE_URL + "/", timeout=5)
        # EitaaBun may not define a root route; any HTTP response below 500
        # confirms that the local process is reachable.
        if response.status_code >= 500:
            raise requests.HTTPError(f"EitaaBun returned HTTP {response.status_code}")
        return {"success": True, "service": "EitaaBun", "status": "online"}
    except requests.RequestException:
        raise HTTPException(503, "سرویس EitaaBun در دسترس نیست")


@router.get("/account")
def account_status():
    saved = _account()
    if not saved or not saved.get("phone"):
        return {"exists": False}
    try:
        result = _request("POST", "/eitaa/auth/loadSession", json={"phone": saved["phone"]})
        if result.get("result") == "ok":
            return {"exists": True, "phone": saved["phone"]}
    except HTTPException:
        pass
    return {"exists": False}


@router.post("/auth/send-code")
def send_code(payload: dict):
    phone = _phone(payload.get("phone"))
    if not phone:
        raise HTTPException(422, "شماره تلفن معتبر نیست")
    result = _request("POST", "/eitaa/auth/sendCode", json={"phone": phone})
    return {"success": True, "phone": phone, "next": result.get("next"), "message": "کد ورود درخواست شد"}


@router.post("/auth/login")
def login(payload: dict):
    phone = _phone(payload.get("phone"))
    code = str(payload.get("code") or "").strip()
    if not phone or not code:
        raise HTTPException(422, "شماره و کد ورود الزامی است")
    try:
        result = _request("POST", "/eitaa/auth/login", json={"phone": phone, "code": code})
    except HTTPException as exc:
        if "PasswordRequired" in str(exc.detail) or "SESSION_PASSWORD_NEEDED" in str(exc.detail):
            return {"success": False, "requires_password": True, "phone": phone, "message": "این حساب رمز دومرحله‌ای دارد"}
        raise
    _save_account(phone)
    return {"success": True, "phone": phone, "user_id": result.get("user_id"), "message": "ورود موفق بود"}


@router.post("/auth/login-password")
def login_password(payload: dict):
    phone = _phone(payload.get("phone"))
    code = str(payload.get("code") or "").strip()
    password = str(payload.get("password") or "")
    if not phone or not code or not password:
        raise HTTPException(422, "شماره، کد و رمز دومرحله‌ای الزامی است")
    result = _request("POST", "/eitaa/auth/loginByPass", json={"phone": phone, "code": code, "password": password})
    _save_account(phone)
    return {"success": True, "phone": phone, "user_id": result.get("user_id"), "message": "ورود موفق بود"}


@router.post("/account/reset")
def reset_account():
    saved = _account()
    if saved and saved.get("phone"):
        try:
            _request("POST", "/eitaa/auth/logout", json={"phone": saved["phone"]})
        except HTTPException:
            pass
    ACCOUNT_FILE.unlink(missing_ok=True)
    return {"success": True}


def _load_recipients(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    first = [str(cell.value or "").strip().lower() for cell in ws[1]]
    phone_names = {"phone", "mobile", "number", "شماره", "شماره تلفن", "موبایل", "تلفن", "شماره همراه"}
    name_names = {"name", "firstname", "first name", "نام", "نام مخاطب"}
    phone_col = next((i + 1 for i, value in enumerate(first) if value in phone_names), None)
    name_col = next((i + 1 for i, value in enumerate(first) if value in name_names), None)
    start_row = 2 if phone_col else 1
    phone_col, name_col = phone_col or 1, name_col or 2
    recipients, seen = [], set()
    for row in range(start_row, ws.max_row + 1):
        phone = _phone(ws.cell(row, phone_col).value)
        name = str(ws.cell(row, name_col).value or "").strip() if name_col <= ws.max_column else ""
        if phone and phone not in seen:
            seen.add(phone)
            recipients.append({"phone": phone, "name": name or phone})
    wb.close()
    return recipients


def _contacts_index(account_phone):
    contacts = _request("GET", "/eitaa/contacts/getAll/" + account_phone).get("contacts", [])
    return {_phone(item.get("phone")): item for item in contacts if item.get("phone")}


def _resolve_contact(account_phone, target, contacts):
    existing = contacts.get(target["phone"])
    if existing and existing.get("id") is not None:
        return {"id": existing["id"], "access_hash": existing.get("access_hash", 0)}
    result = _request("POST", "/eitaa/contacts/add", json={"phone": account_phone, "peerPhone": target["phone"], "firstname": target["name"]})
    if result.get("id") is None:
        raise RuntimeError("مخاطب در ایتا پیدا نشد یا افزودن مخاطب تأیید نشد")
    peer = {"id": result["id"], "access_hash": result.get("acces_hash", result.get("access_hash", 0))}
    contacts[target["phone"]] = {"phone": target["phone"], **peer}
    return peer


def _write_report(job_id, results):
    output = REPORT_DIR / ("eitaa_report_" + job_id + ".xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "نتایج"
    ws.append(["ردیف", "شماره", "نام", "وضعیت", "توضیحات"])
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    for index, item in enumerate(results, 1):
        ws.append([index, item.get("phone", ""), item.get("name", ""), item.get("status", ""), item.get("detail", "")])
    for col, width in {"A": 8, "B": 20, "C": 24, "D": 18, "E": 55}.items():
        ws.column_dimensions[col].width = width
    summary = wb.create_sheet("خلاصه")
    summary.append(["شاخص", "مقدار"])
    summary.append(["کل", len(results)])
    summary.append(["موفق", sum(1 for item in results if item.get("status") == "success")])
    summary.append(["ناموفق", sum(1 for item in results if item.get("status") != "success")])
    wb.save(output)
    return str(output)


def _run_job(job_id, phone, recipients, message, media_path, media_type, delay):
    results = []
    try:
        contacts = _contacts_index(phone)
    except Exception:
        contacts = {}
    with JOBS_LOCK:
        JOBS[job_id]["status"] = "running"
        JOBS[job_id]["started_at"] = datetime.now().isoformat(timespec="seconds")
    for index, target in enumerate(recipients, 1):
        row = {"phone": target["phone"], "name": target["name"], "status": "error", "detail": ""}
        try:
            peer = _resolve_contact(phone, target, contacts)
            payload = {"phone": phone, "id": peer["id"], "access_hash": peer.get("access_hash", 0), "type": "user", "message": message or ""}
            if media_path:
                payload.update({"path": media_path, "media_type": media_type or "file"})
                result = _request("POST", "/eitaa/messages/sendMedia", json=payload)
            else:
                if not message.strip():
                    raise RuntimeError("متن پیام خالی است")
                result = _request("POST", "/eitaa/messages/send", json=payload)
            if result.get("result") != "ok":
                raise RuntimeError("تأیید موفقیت ارسال از سرویس دریافت نشد")
            row["status"], row["detail"] = "success", "پیام ارسال شد"
        except Exception as exc:
            row["detail"] = str(exc)[:500]
        results.append(row)
        with JOBS_LOCK:
            JOBS[job_id]["processed"] = index
            JOBS[job_id]["success"] = sum(1 for item in results if item["status"] == "success")
            JOBS[job_id]["failed"] = len(results) - JOBS[job_id]["success"]
            JOBS[job_id]["results"] = list(results)
        if index < len(recipients):
            time.sleep(delay)
    try:
        report = _write_report(job_id, results)
    except Exception as exc:
        report = None
        with JOBS_LOCK:
            JOBS[job_id]["report_error"] = str(exc)
    with JOBS_LOCK:
        JOBS[job_id]["status"] = "completed"
        JOBS[job_id]["finished_at"] = datetime.now().isoformat(timespec="seconds")
        JOBS[job_id]["report_path"] = report


@router.post("/send")
async def start_send(background_tasks: BackgroundTasks, excel: UploadFile = File(...), message: str = Form(""), media: UploadFile | None = File(None), consent_confirmed: bool = Form(False), delay_seconds: int = Form(3)):
    saved = _account()
    if not saved or not saved.get("phone"):
        raise HTTPException(401, "ابتدا وارد حساب ایتا شوید")
    if not consent_confirmed:
        raise HTTPException(400, "تأیید کنید که مخاطبان اجازه دریافت این پیام را دارند")
    if not message.strip() and not (media and media.filename):
        raise HTTPException(400, "متن پیام یا فایل رسانه‌ای الزامی است")
    if not excel.filename or Path(excel.filename).suffix.lower() not in {".xlsx", ".xlsm"}:
        raise HTTPException(400, "فایل مخاطبان باید Excel با پسوند xlsx یا xlsm باشد")
    job_id = str(uuid.uuid4())
    excel_path = MEDIA_DIR / (job_id + "_" + Path(excel.filename).name)
    media_path = None
    try:
        excel_path.write_bytes(await excel.read())
        recipients = _load_recipients(excel_path)
        if not recipients:
            raise HTTPException(400, "هیچ شماره معتبر و تکرارنشده‌ای در فایل پیدا نشد")
        if len(recipients) > 5000:
            raise HTTPException(400, "حداکثر ۵۰۰۰ مخاطب در هر نوبت مجاز است")
        if media and media.filename:
            suffix = Path(media.filename).suffix.lower()
            if suffix not in {".jpg", ".jpeg", ".png", ".webp", ".gif", ".mp4", ".mov", ".pdf", ".doc", ".docx"}:
                raise HTTPException(400, "فرمت فایل رسانه‌ای پشتیبانی نمی‌شود")
            media_path = MEDIA_DIR / (job_id + "_media" + suffix)
            media_path.write_bytes(await media.read())
        with JOBS_LOCK:
            JOBS[job_id] = {"job_id": job_id, "status": "queued", "total": len(recipients), "processed": 0, "success": 0, "failed": 0, "results": [], "created_at": datetime.now().isoformat(timespec="seconds")}
        delay = max(2, min(int(delay_seconds), 30))
        media_type = "photo" if media_path and media_path.suffix.lower() in {".jpg", ".jpeg", ".png", ".webp", ".gif"} else "file"
        background_tasks.add_task(_run_job, job_id, saved["phone"], recipients, message, str(media_path) if media_path else None, media_type, delay)
        return {"success": True, "job_id": job_id, "total": len(recipients)}
    except Exception:
        excel_path.unlink(missing_ok=True)
        if media_path:
            media_path.unlink(missing_ok=True)
        raise


@router.get("/send/status/{job_id}")
def send_status(job_id: str):
    with JOBS_LOCK:
        job = JOBS.get(job_id)
        if not job:
            raise HTTPException(404, "شناسه ارسال پیدا نشد")
        return {key: value for key, value in job.items() if key != "report_path"}


@router.get("/send/export/{job_id}")
def export_report(job_id: str):
    with JOBS_LOCK:
        job = JOBS.get(job_id)
        report = job.get("report_path") if job else None
    if not report or not Path(report).exists():
        raise HTTPException(404, "گزارش هنوز آماده نیست")
    return FileResponse(report, filename=Path(report).name, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
